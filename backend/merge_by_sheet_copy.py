import pandas as pd
import os
import glob
from collections import defaultdict
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.data_source import StrRef, StrData, StrVal, NumRef, NumData, NumVal, AxDataSource, NumDataSource
from openpyxl.chart.label import DataLabelList

CSV_DIR = "csv_data"
OUTPUT_FILE = "hasil_gabungan_pivot.xlsx"
WORKING_FILE = "hasil_gabungan_pivot_temp.xlsx"

# ── 1. Kumpulkan semua file CSV dan kelompokkan berdasarkan nama sheet ──────
sheet_files = defaultdict(list)
all_csvs = glob.glob(os.path.join(CSV_DIR, "*.csv"))
has_new_format = any("__" in os.path.basename(f) for f in all_csvs)

for filepath in all_csvs:
    basename = os.path.basename(filepath)
    if has_new_format:
        if "__" in basename:
            prefix, sheet_name = basename.replace(".csv", "").split("__", 1)
            sheet_files[sheet_name].append((prefix, filepath))
    else:
        matched = False
        for legacy_prefix in ("SARI_", "NABILLA_"):
            if basename.startswith(legacy_prefix):
                sheet_name = basename[len(legacy_prefix):].replace(".csv", "")
                sheet_files[sheet_name].append((legacy_prefix.rstrip("_"), filepath))
                matched = True
                break
        if not matched:
            parts = basename.replace(".csv", "").split("_", 1)
            if len(parts) == 2:
                sheet_files[parts[1]].append((parts[0], filepath))

# Urutkan file berdasarkan prefix agar data File 1 selalu digabung sebelum File 2
for sheet_name in sheet_files:
    sheet_files[sheet_name].sort(key=lambda x: x[0])

print(f"Ditemukan {len(sheet_files)} sheet unik untuk digabungkan.")

# ── 2. Proses setiap sheet: gabungkan & pivot menjadi tabel ──────────────────
all_code_maps = {}

with pd.ExcelWriter(WORKING_FILE, engine="openpyxl") as writer:

    for sheet_name in sorted(sheet_files.keys()):
        file_entries = sheet_files[sheet_name]
        files = [entry[1] for entry in file_entries]

        dfs = []
        for f in files:
            try:
                df = pd.read_csv(f, encoding="utf-8-sig")
                dfs.append(df)
            except Exception as e:
                print(f"  [SKIP] {f} -> {e}")

        if not dfs:
            continue

        combined = pd.concat(dfs, ignore_index=True)

        required = {"KLASIFIKASI", "JUMLAH KK", "NILAI"}
        safe_sheet = sheet_name[:31]

        # Ambil pemetaan nomor KK -> KODE responden per sheet
        code_map = {}
        if "KODE" in combined.columns:
            for kk_val, grp in combined.groupby("JUMLAH KK"):
                non_empty = grp["KODE"].dropna().astype(str).str.strip()
                non_empty = non_empty[non_empty != ""]
                try:
                    kk_int = int(kk_val)
                except (ValueError, TypeError):
                    kk_int = kk_val
                if len(non_empty) > 0:
                    code_map[kk_int] = non_empty.iloc[0]
                else:
                    code_map[kk_int] = ""

        all_code_maps[safe_sheet] = code_map

        if not required.issubset(combined.columns):
            print(f"  [WARN] {sheet_name}: kolom tidak lengkap -> {combined.columns.tolist()}")
            combined.to_excel(writer, sheet_name=safe_sheet, index=False, startrow=2)
            continue

        combined["NILAI"] = pd.to_numeric(combined["NILAI"], errors="coerce")

        # Deteksi apakah sheet punya sub-klasifikasi
        has_sub = (
            "SUB_KLASIFIKASI" in combined.columns
            and combined["SUB_KLASIFIKASI"].notna().any()
        )

        index_cols = ["KLASIFIKASI", "SUB_KLASIFIKASI"] if has_sub else ["KLASIFIKASI"]

        # ── Urutan asli klasifikasi (sesuai urutan pertama kali muncul di CSV) ──
        order_df = combined[index_cols].drop_duplicates(keep="first")
        if has_sub:
            ordered_index = list(zip(order_df["KLASIFIKASI"], order_df["SUB_KLASIFIKASI"]))
        else:
            ordered_index = list(order_df["KLASIFIKASI"])

        # ── Pivot: gunakan min_count=1 agar sel benar-benar kosong → NaN (bukan 0) ──
        pivot = combined.pivot_table(
            index=index_cols,
            columns="JUMLAH KK",
            values="NILAI",
            aggfunc=lambda x: x.sum(min_count=1),
            dropna=False
        )

        # Urutkan kolom JUMLAH KK secara numerik
        sorted_cols = sorted(pivot.columns, key=lambda x: int(x))
        pivot = pivot[sorted_cols]

        # Reindex sesuai urutan asli (bukan alfabetis hasil pivot_table)
        pivot = pivot.reindex(ordered_index)
        pivot = pivot.reset_index()
        pivot.columns.name = None

        if has_sub:
            no_col = []
            no = 0
            prev_klas = None
            for klas in pivot["KLASIFIKASI"]:
                if klas != prev_klas:
                    no += 1
                    prev_klas = klas
                    no_col.append(no)
                else:
                    no_col.append("")
            pivot.insert(0, "NO", no_col)
        else:
            pivot.insert(0, "NO", range(1, len(pivot) + 1))

        # Ganti NaN dengan "" agar cell tampak kosong
        pivot = pivot.where(pd.notna(pivot), other="")

        # startrow=2:
        # Baris 1: Kode responden
        # Baris 2: JUMLAH KK
        # Baris 3: Header kolom pandas (NO, KLASIFIKASI, 1, 2, ...)
        # Baris 4: Data mulai
        pivot.to_excel(writer, sheet_name=safe_sheet, index=False, startrow=2)

        sumber_list = combined["SUMBER"].unique().tolist()
        print(f"  OK: {sheet_name} ({', '.join(sumber_list)}) -> {len(pivot)} baris, sub={'Ya' if has_sub else 'Tidak'}")


# ── 3. Format dengan openpyxl: tambah header row, merge cells, auto-fit ──────
print("\nMemformat Excel...")

wb = openpyxl.load_workbook(WORKING_FILE)

header1_fill = PatternFill("solid", fgColor="D9E1F2")
header1_font = Font(bold=True)
header2_fill = PatternFill("solid", fgColor="BDD7EE")
header2_font = Font(bold=True)
code_font    = Font(name="Calibri", size=10)

thin   = Side(style="thin", color="999999")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for ws in wb.worksheets:
    n_cols = ws.max_column

    # Deteksi apakah sheet ini punya sub-klasifikasi
    col3_header = ws.cell(row=3, column=3).value
    has_sub_ws  = (col3_header == "SUB_KLASIFIKASI")
    kk_start    = 4 if has_sub_ws else 3
    kk_end      = n_cols

    # ── Tambahkan 15 kolom kosong setelah KK terakhir untuk input manual ────
    EXTRA_KK    = 15
    new_kk_end  = kk_end + EXTRA_KK

    # Kolom JUMLAH dan DLM% di akhir tabel KK
    total_col = new_kk_end + 1
    total_col_letter = get_column_letter(total_col)

    dlm_col = new_kk_end + 2
    dlm_col_letter = get_column_letter(dlm_col)

    # ── Baris 1: Kode Responden ───────────────────────────────────────────
    code_map = all_code_maps.get(ws.title, {})

    for col in range(kk_start, kk_end + 1):
        cell_c = ws.cell(row=1, column=col)
        c_no = ws.cell(row=3, column=col).value
        try:
            c_no_int = int(c_no)
        except (ValueError, TypeError):
            c_no_int = c_no

        code_val = code_map.get(c_no_int, "")
        cell_c.value = code_val
        cell_c.font  = code_font
        cell_c.alignment = Alignment(horizontal="center", vertical="center")
        cell_c.border = border

    for c in range(1, kk_start):
        ws.cell(row=1, column=c).value = ""
        ws.cell(row=1, column=c).border = border
    ws.cell(row=1, column=total_col).border = border
    ws.cell(row=1, column=dlm_col).border = border

    # ── Baris 2: Header JUMLAH KK ─────────────────────────────────────────
    ws.cell(row=2, column=kk_start).value = "JUMLAH KK"
    for col in range(kk_start, new_kk_end + 1):
        c = ws.cell(row=2, column=col)
        c.fill = header1_fill
        c.font = header1_font
        c.border = border

    ws.merge_cells(
        start_row=2, start_column=kk_start,
        end_row=2,   end_column=new_kk_end
    )
    ws.cell(row=2, column=kk_start).alignment = Alignment(
        horizontal="center", vertical="center"
    )

    # ── Merge vertikal Baris 2 & 3 untuk NO, KLASIFIKASI, [SUB] ───────────
    for col in range(1, kk_start):
        ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)
        cell = ws.cell(row=2, column=col)
        cell.fill = header2_fill
        cell.font = header2_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # ── Merge vertikal Baris 2 & 3 untuk JUMLAH dan DLM% ──────────────────
    ws.cell(row=2, column=total_col, value="JUMLAH")
    ws.merge_cells(start_row=2, start_column=total_col, end_row=3, end_column=total_col)
    c_tot = ws.cell(row=2, column=total_col)
    c_tot.fill = header2_fill
    c_tot.font = header2_font
    c_tot.border = border
    c_tot.alignment = Alignment(horizontal="center", vertical="center")

    ws.cell(row=2, column=dlm_col, value="DLM%")
    ws.merge_cells(start_row=2, start_column=dlm_col, end_row=3, end_column=dlm_col)
    c_dlm = ws.cell(row=2, column=dlm_col)
    c_dlm.fill = header2_fill
    c_dlm.font = header2_font
    c_dlm.border = border
    c_dlm.alignment = Alignment(horizontal="center", vertical="center")

    # ── Baris 3: Format nomor KK & Tambah 15 kolom ekstra ─────────────────
    for col in range(kk_start, kk_end + 1):
        c = ws.cell(row=3, column=col)
        c.fill = header2_fill
        c.font = header2_font
        c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center")

    try:
        extra_start_num = int(ws.cell(row=3, column=kk_end).value or (kk_end - kk_start + 1)) + 1
    except (ValueError, TypeError):
        extra_start_num = kk_end - kk_start + 2

    for idx, col in enumerate(range(kk_end + 1, new_kk_end + 1)):
        c1 = ws.cell(row=1, column=col)
        c1.value = ""
        c1.border = border
        c1.alignment = Alignment(horizontal="center", vertical="center")

        c3 = ws.cell(row=3, column=col)
        c3.value = extra_start_num + idx
        c3.fill  = header2_fill
        c3.font  = header2_font
        c3.border = border
        c3.alignment = Alignment(horizontal="center", vertical="center")

    kk_start_letter  = get_column_letter(kk_start)
    kk_end_letter    = get_column_letter(new_kk_end)
    total_kk_count   = new_kk_end - kk_start + 1

    cat_names = []
    totals = []
    dlm_pcts = []
    current_k1 = ""
    data_start = 4

    for r in range(data_start, ws.max_row + 1):
        k1 = ws.cell(row=r, column=2).value
        if k1 is not None and str(k1).strip() != "":
            current_k1 = str(k1).strip()

        if has_sub_ws:
            k2 = ws.cell(row=r, column=3).value
            k2_str = str(k2).strip() if (k2 is not None and str(k2).strip() != "") else ""
            if k2_str:
                cat_names.append(f"{current_k1} - {k2_str}")
            else:
                cat_names.append(current_k1)
        else:
            cat_names.append(current_k1)

        row_total = 0
        for c in range(kk_start, kk_end + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None and v != "":
                try:
                    row_total += float(v)
                except (ValueError, TypeError):
                    pass

        if row_total == int(row_total):
            row_total = int(row_total)
        totals.append(row_total)

        dlm_val = (row_total / total_kk_count) if total_kk_count > 0 else 0.0
        dlm_pcts.append(dlm_val)

        cell_tot = ws.cell(row=r, column=total_col)
        cell_tot.value = f"=SUM({kk_start_letter}{r}:{kk_end_letter}{r})"
        cell_tot.font = Font(bold=True)
        cell_tot.border = border
        cell_tot.alignment = Alignment(horizontal="center", vertical="center")

        cell_dlm = ws.cell(row=r, column=dlm_col)
        cell_dlm.value = f"={total_col_letter}{r}/{total_kk_count}"
        cell_dlm.number_format = "0%"
        cell_dlm.font = Font(bold=True)
        cell_dlm.border = border
        cell_dlm.alignment = Alignment(horizontal="center", vertical="center")

    # ── Tambah baris TOTAL di bawah baris data terakhir ───────────────────
    total_row_idx = ws.max_row + 1
    total_fill = PatternFill("solid", fgColor="FFC000")
    total_font = Font(bold=True)

    label_end = kk_start - 1
    for c in range(1, label_end + 1):
        cell_lbl = ws.cell(row=total_row_idx, column=c)
        cell_lbl.fill   = total_fill
        cell_lbl.border = border
    ws.merge_cells(
        start_row=total_row_idx, start_column=1,
        end_row=total_row_idx,   end_column=label_end
    )
    lbl = ws.cell(row=total_row_idx, column=1)
    lbl.value     = "TOTAL"
    lbl.font      = total_font
    lbl.alignment = Alignment(horizontal="center", vertical="center")

    for c in range(kk_start, new_kk_end + 1):
        ct = ws.cell(row=total_row_idx, column=c)
        ct.fill   = total_fill
        ct.border = border

    ct_total = ws.cell(row=total_row_idx, column=total_col)
    ct_total.value     = f"=SUM({total_col_letter}{data_start}:{total_col_letter}{total_row_idx - 1})"
    ct_total.font      = total_font
    ct_total.fill      = total_fill
    ct_total.border    = border
    ct_total.alignment = Alignment(horizontal="center", vertical="center")

    ct_dlm = ws.cell(row=total_row_idx, column=dlm_col)
    ct_dlm.value        = f"=AVERAGE({dlm_col_letter}{data_start}:{dlm_col_letter}{total_row_idx - 1})"
    ct_dlm.number_format = "0%"
    ct_dlm.font         = total_font
    ct_dlm.fill         = total_fill
    ct_dlm.border       = border
    ct_dlm.alignment    = Alignment(horizontal="center", vertical="center")

    for row in range(data_start, total_row_idx):
        for col in range(1, dlm_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = border
            if col != total_col and col != dlm_col:
                cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row, column=2).alignment = Alignment(
            horizontal="left", vertical="center"
        )
        if has_sub_ws:
            ws.cell(row=row, column=3).alignment = Alignment(
                horizontal="left", vertical="center"
            )

    # ── Merge cells NO & KLASIFIKASI untuk sheet dengan sub-klasifikasi ────
    if has_sub_ws:
        groups = []
        current_start = None

        for excel_row in range(data_start, total_row_idx):
            no_val = ws.cell(row=excel_row, column=1).value
            if no_val is not None and no_val != "":
                if current_start is not None:
                    groups.append((current_start, excel_row - 1))
                current_start = excel_row

        if current_start is not None:
            groups.append((current_start, total_row_idx - 1))

        for start_r, end_r in groups:
            if end_r > start_r:
                ws.merge_cells(
                    start_row=start_r, start_column=1,
                    end_row=end_r,   end_column=1
                )
                ws.merge_cells(
                    start_row=start_r, start_column=2,
                    end_row=end_r,   end_column=2
                )
                ws.cell(start_r, 1).alignment = Alignment(
                    horizontal="center", vertical="center"
                )
                ws.cell(start_r, 2).alignment = Alignment(
                    horizontal="left", vertical="center"
                )

    # ── Auto-fit lebar kolom berdasarkan konten ────────────────────────────
    for col_idx in range(1, dlm_col + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row in range(1, ws.max_row + 1):
            cell = ws.cell(row=row, column=col_idx)
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 6), 30)

    ws.column_dimensions[total_col_letter].width = 12
    ws.column_dimensions[dlm_col_letter].width = 12
    for col in range(kk_end + 1, new_kk_end + 1):
        ws.column_dimensions[get_column_letter(col)].width = 7
    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 16

    # ── Freeze Panes: kolom NO, KLASIFIKASI (dan SUB) serta baris header ───
    freeze_col_letter = get_column_letter(kk_start)
    ws.freeze_panes = f"{freeze_col_letter}{data_start}"

    # ── Tambah Diagram Batang di sebelah kanan tabel ────────────────────────
    if ws.max_row >= 4 and cat_names:
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = ws.title.replace("_", " ")

        chart.x_axis.delete = False
        chart.y_axis.delete = False
        chart.x_axis.axPos = "b"
        chart.y_axis.axPos = "l"
        chart.x_axis.tickLblPos = "nextTo"
        chart.y_axis.tickLblPos = "nextTo"
        chart.y_axis.majorGridlines = ChartLines()
        chart.y_axis.number_format = "0%"

        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True
        chart.dataLabels.numFmt = "0%"

        chart.legend = None
        chart.width = max(16, min(len(cat_names) * 2.2, 35))
        chart.height = 11

        chart_data_end = total_row_idx - 1
        data_ref = Reference(ws, min_col=dlm_col, min_row=3, max_row=chart_data_end)
        chart.add_data(data_ref, titles_from_data=True)

        str_pts = [StrVal(idx=i, v=cat_names[i]) for i in range(len(cat_names))]
        str_data = StrData(pt=str_pts, ptCount=len(str_pts))
        safe_title = ws.title
        str_ref = StrRef(f"'{safe_title}'!$B$4:$B${chart_data_end}", strCache=str_data)
        chart.series[0].cat = AxDataSource(strRef=str_ref)

        num_pts = [NumVal(idx=i, v=float(dlm_pcts[i])) for i in range(len(dlm_pcts))]
        num_data = NumData(pt=num_pts, ptCount=len(num_pts))
        num_ref = NumRef(f"'{safe_title}'!${dlm_col_letter}$4:${dlm_col_letter}${chart_data_end}", numCache=num_data)
        chart.series[0].val = NumDataSource(numRef=num_ref)

        chart_anchor = f"{get_column_letter(dlm_col + 2)}2"
        ws.add_chart(chart, chart_anchor)

# ── Simpan Hasil Workbook ──────────────────────────────────────────────────
wb.save(OUTPUT_FILE)
print(f"Selesai! Output tersimpan di: {OUTPUT_FILE}")

# Bersihkan file kerja sementara jika ada
if os.path.exists(WORKING_FILE) and WORKING_FILE != OUTPUT_FILE:
    try:
        os.remove(WORKING_FILE)
    except Exception:
        pass
